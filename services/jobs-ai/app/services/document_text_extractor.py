from __future__ import annotations

from io import BytesIO

from docx import Document
from pypdf import PdfReader

from app.schemas.resume_extraction import ResumeTextExtractionResponse

PDF_MIME_TYPE = "application/pdf"
DOCX_MIME_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
SUPPORTED_MIME_TYPES = {PDF_MIME_TYPE, DOCX_MIME_TYPE}


def extract_document_text(
    content: bytes,
    detected_mime_type: str,
) -> ResumeTextExtractionResponse:
    if detected_mime_type not in SUPPORTED_MIME_TYPES:
        return build_response(
            status="unsupported_type",
            detected_mime_type=detected_mime_type,
            extraction_method="none",
            warnings=["Unsupported resume file type."],
        )

    if not content:
        return build_response(
            status="empty_document",
            detected_mime_type=detected_mime_type,
            extraction_method="none",
            warnings=["Resume document is empty."],
        )

    if detected_mime_type == PDF_MIME_TYPE:
        return extract_pdf_text(content, detected_mime_type)

    return extract_docx_text(content, detected_mime_type)


def extract_pdf_text(content: bytes, detected_mime_type: str) -> ResumeTextExtractionResponse:
    try:
        reader = PdfReader(BytesIO(content))
    except Exception:
        return build_response(
            status="corrupt_document",
            detected_mime_type=detected_mime_type,
            extraction_method="pypdf",
            warnings=["PDF could not be read."],
        )

    if reader.is_encrypted:
        return build_response(
            status="encrypted_document",
            detected_mime_type=detected_mime_type,
            extraction_method="pypdf",
            warnings=["PDF is encrypted or password protected."],
            page_count=len(reader.pages),
        )

    page_count = len(reader.pages)

    if page_count == 0:
        return build_response(
            status="empty_document",
            detected_mime_type=detected_mime_type,
            extraction_method="pypdf",
            warnings=["PDF contains no pages."],
            page_count=0,
        )

    try:
        text = "\n".join(page.extract_text() or "" for page in reader.pages).strip()
    except Exception:
        return build_response(
            status="extraction_failed",
            detected_mime_type=detected_mime_type,
            extraction_method="pypdf",
            warnings=["PDF text extraction failed."],
            page_count=page_count,
        )

    if not text:
        return build_response(
            status="ocr_required",
            detected_mime_type=detected_mime_type,
            extraction_method="pypdf",
            warnings=["No embedded PDF text was found; OCR is required."],
            page_count=page_count,
            requires_ocr=True,
        )

    return build_response(
        status="text_extracted",
        text=text,
        detected_mime_type=detected_mime_type,
        extraction_method="pypdf",
        page_count=page_count,
    )


def extract_docx_text(content: bytes, detected_mime_type: str) -> ResumeTextExtractionResponse:
    try:
        document = Document(BytesIO(content))
    except Exception:
        return build_response(
            status="corrupt_document",
            detected_mime_type=detected_mime_type,
            extraction_method="python-docx",
            warnings=["DOCX could not be read."],
        )

    parts: list[str] = []
    parts.extend(paragraph.text for paragraph in document.paragraphs if paragraph.text.strip())

    for table in document.tables:
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            if row_text:
                parts.append(row_text)

    text = "\n".join(parts).strip()

    if not text:
        return build_response(
            status="empty_document",
            detected_mime_type=detected_mime_type,
            extraction_method="python-docx",
            warnings=["DOCX contains no extractable text."],
        )

    return build_response(
        status="text_extracted",
        text=text,
        detected_mime_type=detected_mime_type,
        extraction_method="python-docx",
    )


def build_response(
    *,
    status: str,
    detected_mime_type: str,
    extraction_method: str,
    text: str = "",
    page_count: int | None = None,
    warnings: list[str] | None = None,
    requires_ocr: bool = False,
) -> ResumeTextExtractionResponse:
    return ResumeTextExtractionResponse(
        status=status,
        text=text,
        pageCount=page_count,
        detectedMimeType=detected_mime_type,
        characterCount=len(text),
        extractionMethod=extraction_method,
        warnings=warnings or [],
        requiresOcr=requires_ocr,
    )
